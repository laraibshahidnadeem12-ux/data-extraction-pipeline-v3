"""Persistence layer: Excel, Google Sheets, CSV, JSON record store."""
from __future__ import annotations

import csv
import io
import json
import os
import threading

from openpyxl import Workbook, load_workbook

from .models import ExtractedRecord

_lock = threading.Lock()
HEADERS = ExtractedRecord.COLUMNS()


def _ensure_excel(path: str) -> None:
    if os.path.exists(path):
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "ExtractedData"
    ws.append(HEADERS)
    wb.save(path)


def append_to_excel(records: list[ExtractedRecord], path: str) -> str:
    """Append records to an .xlsx workbook (creates it with headers first)."""
    path = os.path.abspath(path)
    with _lock:
        _ensure_excel(path)
        wb = load_workbook(path)
        ws = wb["ExtractedData"]
        for record in records:
            ws.append(record.to_row())
        wb.save(path)
    return path


def load_records(records_path: str) -> list[dict]:
    if not os.path.exists(records_path):
        return []
    with open(records_path, "r", encoding="utf-8") as fh:
        return json.load(fh)


def save_records(records: list[dict], records_path: str) -> None:
    with _lock:
        os.makedirs(os.path.dirname(os.path.abspath(records_path)), exist_ok=True)
        with open(records_path, "w", encoding="utf-8") as fh:
            json.dump(records, fh, ensure_ascii=False, indent=2)


def records_to_csv(records: list[dict]) -> str:
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=HEADERS, extrasaction="ignore")
    writer.writeheader()
    for record in records:
        writer.writerow(
            {
                col: json.dumps(record.get(col), ensure_ascii=False)
                if isinstance(record.get(col), (dict, list))
                else record.get(col, "")
                for col in HEADERS
            }
        )
    return buf.getvalue()


def sync_to_sheets(records: list[dict], config: dict) -> bool:
    """Append a batch of records to a Google Sheet via a service account."""
    sheets_cfg = config.get("google_sheets", {})
    if not sheets_cfg.get("enabled"):
        return False
    service_account = sheets_cfg.get("service_account_json", "service_account.json")
    spreadsheet_id = sheets_cfg.get("spreadsheet_id", "")
    tab = sheets_cfg.get("sheet_tab", "ExtractedData")
    if not spreadsheet_id or not os.path.exists(service_account):
        print(f"[sheets] skipped: enabled but no spreadsheet_id/service account at {service_account}")
        return False
    try:
        import gspread
        gc = gspread.service_account(filename=service_account)
        sheet = gc.open_by_key(spreadsheet_id)
        ws = sheet.worksheet(tab) if tab in [s.title for s in sheet.worksheets()] else sheet.add_worksheet(tab, 100, len(HEADERS))
        if not ws.get_all_values():
            ws.append_row(HEADERS)
        rows = [[record.get(col, "") for col in HEADERS] for record in records]
        ws.append_rows(rows)
        return True
    except Exception as exc:  # keep the pipeline alive if Sheets fails
        print(f"[sheets] sync failed: {exc}")
        return False